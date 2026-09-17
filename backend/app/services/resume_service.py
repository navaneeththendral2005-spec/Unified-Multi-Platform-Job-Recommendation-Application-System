import re

from pypdf import PdfReader
from docx import Document


# ============================================================
# TEXT CLEANING
# ============================================================

def clean_extracted_text(
    text: str
) -> str:
    """
    Clean and normalize extracted resume text.
    """

    if not text:
        return ""

    # Normalize line endings.

    text = text.replace(
        "\r\n",
        "\n"
    )

    text = text.replace(
        "\r",
        "\n"
    )

    # Remove excessive spaces.

    text = re.sub(
        r"[ \t]+",
        " ",
        text
    )

    # Remove excessive blank lines.

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text
    )

    return text.strip()


# ============================================================
# PDF TEXT EXTRACTION
# ============================================================

def extract_text_from_pdf(
    file_path: str
) -> str:
    """
    Extract text content from a PDF resume.
    """

    try:

        reader = PdfReader(file_path)

        extracted_text = []

        for page in reader.pages:

            page_text = page.extract_text()

            if page_text:

                extracted_text.append(
                    page_text
                )

        text = "\n".join(
            extracted_text
        )

        return clean_extracted_text(
            text
        )

    except Exception as error:

        raise ValueError(
            f"Failed to extract text from PDF: {str(error)}"
        )


# ============================================================
# DOCX TEXT EXTRACTION
# ============================================================

def extract_text_from_docx(
    file_path: str
) -> str:
    """
    Extract text content from a DOCX resume.
    """

    try:

        document = Document(
            file_path
        )

        extracted_text = []

        # ----------------------------------------------------
        # EXTRACT PARAGRAPHS
        # ----------------------------------------------------

        for paragraph in document.paragraphs:

            paragraph_text = (
                paragraph.text.strip()
            )

            if paragraph_text:

                extracted_text.append(
                    paragraph_text
                )

        # ----------------------------------------------------
        # EXTRACT TABLE CONTENT
        # ----------------------------------------------------

        for table in document.tables:

            for row in table.rows:

                row_cells = []

                for cell in row.cells:

                    cell_text = (
                        cell.text.strip()
                    )

                    if cell_text:

                        row_cells.append(
                            cell_text
                        )

                if row_cells:

                    extracted_text.append(
                        " | ".join(row_cells)
                    )

        text = "\n".join(
            extracted_text
        )

        return clean_extracted_text(
            text
        )

    except Exception as error:

        raise ValueError(
            f"Failed to extract text from DOCX: {str(error)}"
        )


# ============================================================
# TEXT QUALITY VALIDATION
# ============================================================

MINIMUM_RESUME_CHARACTERS = 100

MINIMUM_RESUME_WORDS = 20


def validate_resume_text(
    text: str
) -> dict:
    """
    Validate whether extracted resume text
    contains enough meaningful information
    for analysis.
    """

    cleaned_text = clean_extracted_text(
        text
    )

    character_count = len(
        cleaned_text
    )

    words = re.findall(
        r"\b\w+\b",
        cleaned_text
    )

    word_count = len(
        words
    )

    lines = [

        line

        for line in cleaned_text.splitlines()

        if line.strip()
    ]

    line_count = len(
        lines
    )

    # --------------------------------------------------------
    # EMPTY RESUME
    # --------------------------------------------------------

    if not cleaned_text:

        return {

            "is_valid": False,

            "quality_level": "empty",

            "message": (
                "No readable text could be extracted "
                "from the resume."
            ),

            "character_count": 0,

            "word_count": 0,

            "line_count": 0
        }

    # --------------------------------------------------------
    # VERY WEAK EXTRACTION
    # --------------------------------------------------------

    if (
        character_count < MINIMUM_RESUME_CHARACTERS
        or word_count < MINIMUM_RESUME_WORDS
    ):

        return {

            "is_valid": False,

            "quality_level": "insufficient",

            "message": (
                "The extracted resume text is too short "
                "for reliable analysis."
            ),

            "character_count": character_count,

            "word_count": word_count,

            "line_count": line_count
        }

    # --------------------------------------------------------
    # VALID RESUME TEXT
    # --------------------------------------------------------

    return {

        "is_valid": True,

        "quality_level": "good",

        "message": (
            "Resume text was successfully extracted "
            "and is ready for analysis."
        ),

        "character_count": character_count,

        "word_count": word_count,

        "line_count": line_count
    }